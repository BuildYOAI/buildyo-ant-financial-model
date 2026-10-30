"""Build ANT_Financial_Model.xlsx from data/historicals.json, data/forecast.json and
model/assumptions.json (all read-only inputs).

Run:  python3 scripts/build_workbook.py
Refresh forecasts first with `node model/run.mjs` if model/assumptions.json changed.
All amounts in $ millions.
"""
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent.parent
HIST = json.loads((ROOT / "data" / "historicals.json").read_text())
FC = json.loads((ROOT / "data" / "forecast.json").read_text())
ASSUME = json.loads((ROOT / "model" / "assumptions.json").read_text())
OUT = ROOT / "ANT_Financial_Model.xlsx"

# ---------------------------------------------------------------- styles
FONT = "Arial"
NUM = '#,##0;(#,##0);"–"'
NUM1 = '#,##0.0;(#,##0.0);"–"'
PCT0, PCT1 = "0%", "0.0%"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_ITAL = Font(name=FONT, size=10, italic=True)
F_DERIVED = Font(name=FONT, size=10, italic=True, color="595959")
F_DERIVED_B = Font(name=FONT, size=10, italic=True, bold=True, color="595959")
F_NE = Font(name=FONT, size=10, color="A6A6A6")
F_LINK = Font(name=FONT, size=10, color="008000")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_HDR = Font(name=FONT, size=10, bold=True)
FILL_PRESS = PatternFill("solid", fgColor="FCE4D6")   # light orange
FILL_EST = PatternFill("solid", fgColor="FFF2CC")     # light yellow
FILL_FC = PatternFill("solid", fgColor="DDEBF7")      # light blue
FILL_HIST_HDR = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="808080")
B_TOP = Border(top=THIN)
B_BOTTOM = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")
CENTER = Alignment(horizontal="center")

SCEN = {
    "high": ("Amazon Path", "The Amazon Path (High Momentum)"),
    "decel": ("Cisco Path", "The Cisco Path (Deceleration)"),
    "decline": ("AOL Path", "The AOL Path (Decline)"),
}
S1_URL = HIST["meta"]["s1"]


def q(sheet):
    return f"'{sheet}'"


def style_val(c, flag=None, fmt=NUM, bold=False, italic=False):
    """Apply number format + source-flag styling to a value cell."""
    c.number_format = fmt
    if flag == "derived":
        c.font = F_DERIVED_B if bold else F_DERIVED
    elif flag == "n/e":
        c.font = F_NE
        c.alignment = RIGHT
    else:
        c.font = F_BOLD if bold else (F_ITAL if italic else F_BASE)
    if flag == "press":
        c.fill = FILL_PRESS
    elif flag == "est":
        c.fill = FILL_EST


def header_row(ws, row, labels, start_col=2, fill=FILL_HIST_HDR, first="$ millions"):
    c = ws.cell(row, 1, first)
    c.font = F_NOTE
    for j, lab in enumerate(labels):
        c = ws.cell(row, start_col + j, lab)
        c.font = F_HDR
        c.alignment = CENTER
        c.fill = fill if not isinstance(fill, list) else fill[j]
        c.border = B_BOTTOM


def title(ws, text, sub=None):
    ws.cell(1, 1, text).font = F_TITLE
    if sub:
        ws.cell(2, 1, sub).font = F_NOTE


def widths(ws, first=46, rest=11, ncols=20):
    ws.column_dimensions["A"].width = first
    for j in range(2, ncols + 2):
        ws.column_dimensions[COL(j)].width = rest


def label(ws, row, text, bold=False, italic=False, indent=0):
    c = ws.cell(row, 1, text)
    c.font = F_BOLD if bold else (F_ITAL if italic else F_BASE)
    if indent:
        c.alignment = Alignment(indent=indent)
    return c


def margin_formula(num_ref, den_ref):
    return f'=IF({den_ref}=0,"",{num_ref}/{den_ref})'


# ---------------------------------------------------------------- README
def build_readme(wb):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    title(ws, "Anthropic (ANT) — S-1 Financial Model")
    r = 3
    rows = [
        ("Generated", f"{date.today().isoformat()} (forecast engine output timestamp: {FC.get('generated', 'n/a')})"),
        ("Source (S-1)", S1_URL),
        ("Units", "US$ millions unless stated (headcount in people; shares in millions; IPO price in $/share)"),
    ]
    for k, v in rows:
        ws.cell(r, 1, k).font = F_BOLD
        c = ws.cell(r, 2, v)
        c.font = F_BASE
        if k == "Source (S-1)":
            c.hyperlink = v
            c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        r += 1

    r += 1
    ws.cell(r, 1, "Color legend").font = F_SUB
    r += 1
    legend = [
        ("1,234", F_BASE, None, "S-1 stated (default black text)"),
        ("1,234", F_DERIVED, None, "Derived — arithmetic on S-1 figures (dark gray italic), e.g. H2'25 = FY25 − H1'25"),
        ("1,234", F_BASE, FILL_PRESS, "Press-reported (Oct 14–15 2026 investor presentation via Bloomberg / The Information / CNBC / FT)"),
        ("1,234", F_BASE, FILL_EST, "Model estimate (assumption; see Hist IS (Qtr) notes)"),
        ("n/e", F_NE, None, "Not estimable — excluded from totals"),
        ("Q1'27", F_HDR, FILL_FC, "Forecast column header (engine output, not a prediction)"),
        ("1,234", F_LINK, None, "Green text = formula linking to another sheet in this workbook"),
        ("1,234", F_BOLD, None, "Bold = total / subtotal (historical subtotals are live Excel formulas)"),
        ("1,234", F_ITAL, None, "Italic = memo line (not part of the totals above it)"),
    ]
    for sample, font, fill, desc in legend:
        c = ws.cell(r, 1, sample)
        c.font = font
        c.alignment = RIGHT
        if fill:
            c.fill = fill
        ws.cell(r, 2, desc).font = F_BASE
        r += 1

    r += 1
    ws.cell(r, 1, "Scenarios").font = F_SUB
    r += 1
    for key, (_, name) in SCEN.items():
        s = ASSUME["scenarios"][key]
        ipo = FC["scenarios"][key]["ipo"]
        text = (f"{s.get('_bud', '')}  Model IPO: {ipo['quarter']} at ${ipo['price']:.2f}/share "
                f"(primary raise ${ipo['primary']:,.0f}M + Google placement ${ipo['googlePlacement']:,.0f}M).")
        ws.cell(r, 1, name).font = F_BOLD
        ws.cell(r, 1).alignment = WRAP
        c = ws.cell(r, 2, text)
        c.font = F_BASE
        c.alignment = WRAP
        ws.row_dimensions[r].height = 45
        r += 1

    r += 1
    ws.cell(r, 1, "Caveats").font = F_SUB
    r += 1
    caveats = [
        "Forecasts are model outputs, not predictions. They show what the stated assumptions imply, nothing more.",
        "Q3'26 is press-reported (revenue, compute, adjusted operating income) plus flagged model estimates; it is not in the S-1. "
        "Fair-value marks on the convertible notes for Q3'26 are not estimable and are excluded.",
        "The dotcom analogs (Amazon, Cisco, AOL) describe the shape of a trajectory, not its scale.",
        "Units are US$ millions throughout unless a row says otherwise.",
        "The forecast cash flow is on a cash + short-term-investments basis (purchases/sales of ST investments are not shown as investing flows).",
        "Forecasts are generated by model/engine.js. To change them: edit model/assumptions.json, run `node model/run.mjs`, "
        "then rerun `python3 scripts/build_workbook.py`.",
        "Historical subtotals are Excel formulas over the S-1 line items; forecast sheets hold engine values (not formulas).",
    ]
    for t in caveats:
        c = ws.cell(r, 2, "• " + t)
        c.font = F_BASE
        c.alignment = WRAP
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    ws.cell(r, 1, "Sheets").font = F_SUB
    r += 1
    for name, desc in [
        ("Hist IS (Qtr)", "Quarterly income statement Q1'25–Q2'26 (S-1) + Q3'26 (press/estimate), margins"),
        ("Hist IS (Annual)", "FY23–FY25, H1'25, H1'26 income statement + functional opex view (S-1 Note 8)"),
        ("Hist BS", "Balance sheets Dec'24, Dec'25, Jun'26 with check row"),
        ("Hist CF", "Cash flow FY23–FY25, H1'25, H2'25 (derived), H1'26; free cash flow"),
        ("Amazon Path / Cisco Path / AOL Path", "Forecast IS, BS, CF and key ratios: Q3'26–Q4'27 quarterly + FY2026–FY2030"),
        ("Compute Staircase", "Quarterly compute cost by contract, per scenario, with charts"),
        ("Scenario Compare", "Annual 2026–2030 key metrics across scenarios, with charts"),
        ("Assumptions", "Flattened model/assumptions.json and contracts table"),
    ]:
        ws.cell(r, 1, name).font = F_BOLD
        ws.cell(r, 2, desc).font = F_BASE
        r += 1


# ---------------------------------------------------------------- Hist IS (Qtr)
QIS_ROWS = [  # key, label, kind ; kind: v=value, f=formula
    ("revenue", "Revenue", "v"),
    ("compute", "Compute & infrastructure", "v"),
    ("personnel", "Personnel", "v"),
    ("bartz", "Bartz settlement", "v"),
    ("sales_marketing", "Sales, marketing & partnerships", "v"),
    ("other_opex", "Other operating & administrative", "v"),
    ("total_opex", "Total operating expenses", "f"),
    ("operating_income", "Operating income (loss)", "f"),
    ("interest_income", "Interest income", "v"),
    ("fv_change_notes", "Change in fair value of convertible notes & SAFEs", "v"),
    ("other_income", "Other income (expense), net", "v"),
    ("pretax_income", "Pre-tax income (loss)", "f"),
    ("tax", "Provision for income taxes", "v"),
    ("net_income", "Net income (loss)", "f"),
]


def is_formulas(rowmap, c):
    """Subtotal formulas for an income statement column (shared by quarterly/annual)."""
    R = lambda k: f"{c}{rowmap[k]}"
    return {
        "total_opex": f"=SUM({R('compute')}:{R('other_opex')})",
        "operating_income": f"={R('revenue')}-{R('total_opex')}",
        "pretax_income": f"=SUM({R('operating_income')}:{R('other_income')})",  # SUM skips n/e text
        "net_income": f"={R('pretax_income')}-{R('tax')}",
    }


def build_hist_is_q(wb):
    ws = wb.create_sheet("Hist IS (Qtr)")
    d = HIST["quarterly_is"]
    periods, lines = d["periods"], d["lines"]
    title(ws, "Historical income statement — quarterly", "Q1'25–Q2'26 from S-1 MD&A quarterly table (unaudited); Q3'26 press-reported + model estimates (see notes below)")
    hr = 4
    header_row(ws, hr, periods)
    r = hr + 1
    rowmap = {}
    for key, lab, kind in QIS_ROWS:
        rowmap[key] = r
        label(ws, r, lab, bold=(kind == "f"))
        r += 1
    r += 1
    label(ws, r, "Memo (non-GAAP)", bold=True); r += 1
    for key, lab in [("sbc", "Stock-based compensation"), ("noncash_charity", "Non-cash donation matching"),
                     ("adj_operating_income", "Adjusted operating income (op. income + SBC + donations + Bartz)"),
                     ("adj_net_income", "Adjusted net income (S-1 non-GAAP)")]:
        rowmap[key] = r
        label(ws, r, lab, italic=True)
        r += 1
    r += 1
    label(ws, r, "Margins", bold=True); r += 1
    for key, lab in [("op_margin", "Operating margin"), ("adj_op_margin", "Adjusted operating margin"),
                     ("compute_pct", "Compute % of revenue")]:
        rowmap[key] = r
        label(ws, r, lab, italic=True)
        r += 1

    for j, p in enumerate(periods):
        col = 2 + j
        c = COL(col)
        f = is_formulas(rowmap, c)
        f["adj_operating_income"] = (f"={c}{rowmap['operating_income']}+{c}{rowmap['sbc']}"
                                     f"+{c}{rowmap['noncash_charity']}+{c}{rowmap['bartz']}")
        for key in [k for k, _, _ in QIS_ROWS] + ["sbc", "noncash_charity", "adj_operating_income", "adj_net_income"]:
            val, flag = lines[key]["values"][j], lines[key]["flags"][j]
            cell = ws.cell(rowmap[key], col)
            if key in f:
                cell.value = f[key]
            elif flag == "n/e" or val is None:
                cell.value = "n/e"
            else:
                cell.value = val
            is_total = key in ("total_opex", "operating_income", "pretax_income", "net_income")
            style_val(cell, flag, bold=is_total, italic=key in ("sbc", "noncash_charity", "adj_operating_income", "adj_net_income"))
            if is_total:
                cell.border = B_TOP
        rev = f"{c}{rowmap['revenue']}"
        for key, num in [("op_margin", "operating_income"), ("adj_op_margin", "adj_operating_income"), ("compute_pct", "compute")]:
            cell = ws.cell(rowmap[key], col, margin_formula(f"{c}{rowmap[num]}", rev))
            style_val(cell, None, fmt=PCT1, italic=True)

    r += 1
    ws.cell(r, 1, "Q3'26 sources and assumptions").font = F_SUB
    r += 1
    for k, v in HIST["meta"]["q3_assumptions"].items():
        ws.cell(r, 1, k).font = F_NOTE
        ws.cell(r, 2, v).font = F_NOTE
        r += 1
    ws.cell(r, 1, "Q3'26 donation-matching charge ($2.0B) is stated in S-1 Note 17.").font = F_NOTE
    widths(ws, 58, 11, len(periods))
    ws.freeze_panes = ws.cell(hr + 1, 2)
    return rowmap


# ---------------------------------------------------------------- Hist IS (Annual)
def build_hist_is_a(wb):
    ws = wb.create_sheet("Hist IS (Annual)")
    d = HIST["annual_is"]
    periods = d["periods"]
    title(ws, "Historical income statement — annual and six-month", "S-1 consolidated statements of operations (pp. F-6 to F-12)")
    hr = 4
    header_row(ws, hr, periods)
    r = hr + 1
    rowmap = {}
    for key, lab, kind in QIS_ROWS:
        rowmap[key] = r
        label(ws, r, lab, bold=(kind == "f"))
        r += 1
    rowmap["op_margin"] = r
    label(ws, r, "Operating margin", italic=True)
    r += 1
    rowmap["compute_pct"] = r
    label(ws, r, "Compute % of revenue", italic=True)
    r += 2

    ws.cell(r, 1, "Functional operating expenses (S-1 Note 8)").font = F_SUB
    r += 1
    header_row(ws, r, periods)
    r += 1
    fmap = {}
    for key, lab in [("cost_of_revenue", "Cost of revenue"), ("research_development", "Research & development"),
                     ("sales_marketing", "Sales & marketing"), ("bartz", "Bartz settlement"),
                     ("general_admin", "General & administrative"), ("total_opex", "Total operating expenses"),
                     ("check", "Check: functional total − income statement total (should be 0)"),
                     ("gross_profit", "Gross profit (revenue − cost of revenue)"), ("gross_margin", "Gross margin")]:
        fmap[key] = r
        label(ws, r, lab, bold=key in ("total_opex", "gross_profit"), italic=key in ("check", "gross_margin"))
        r += 1

    for j in range(len(periods)):
        col = 2 + j
        c = COL(col)
        f = is_formulas(rowmap, c)
        for key, _, _ in QIS_ROWS:
            cell = ws.cell(rowmap[key], col, f.get(key, d[key][j]))
            tot = key in f
            style_val(cell, "s1", bold=tot)
            if tot:
                cell.border = B_TOP
        rev = f"{c}{rowmap['revenue']}"
        style_val(ws.cell(rowmap["op_margin"], col, margin_formula(f"{c}{rowmap['operating_income']}", rev)), fmt=PCT1, italic=True)
        style_val(ws.cell(rowmap["compute_pct"], col, margin_formula(f"{c}{rowmap['compute']}", rev)), fmt=PCT1, italic=True)
        fo = HIST["functional_opex"]
        for key in ("cost_of_revenue", "research_development", "sales_marketing", "bartz", "general_admin"):
            style_val(ws.cell(fmap[key], col, fo[key][j]), "s1")
        cell = ws.cell(fmap["total_opex"], col, f"=SUM({c}{fmap['cost_of_revenue']}:{c}{fmap['general_admin']})")
        style_val(cell, bold=True); cell.border = B_TOP
        style_val(ws.cell(fmap["check"], col, f"={c}{fmap['total_opex']}-{c}{rowmap['total_opex']}"), italic=True)
        style_val(ws.cell(fmap["gross_profit"], col, f"={rev}-{c}{fmap['cost_of_revenue']}"), bold=True)
        style_val(ws.cell(fmap["gross_margin"], col, margin_formula(f"{c}{fmap['gross_profit']}", rev)), fmt=PCT1, italic=True)
    widths(ws, 58, 11, len(periods))
    ws.freeze_panes = ws.cell(hr + 1, 2)
    return rowmap, fmap


# ---------------------------------------------------------------- Hist BS
BS_LAYOUT = [  # key, label, type ; type v=value, or formula spec
    ("cash", "Cash and cash equivalents", "v"),
    ("st_investments", "Short-term investments", "v"),
    ("accounts_receivable", "Accounts receivable, net", "v"),
    ("prepaid_other_current", "Prepaid expenses & other current assets", "v"),
    ("total_current_assets", "Total current assets", ("sum", "cash", "prepaid_other_current")),
    ("restricted_cash", "Restricted cash, non-current", "v"),
    ("ppe_net", "Property & equipment, net", "v"),
    ("rou_assets", "Operating lease right-of-use assets", "v"),
    ("goodwill", "Goodwill", "v"),
    ("other_noncurrent_assets", "Other non-current assets", "v"),
    ("total_assets", "Total assets", ("sum", "total_current_assets", "other_noncurrent_assets")),
    None,
    ("ap_accrued", "Accounts payable & accrued liabilities", "v"),
    ("lease_liab_current", "Operating lease liabilities, current", "v"),
    ("deferred_revenue", "Deferred revenue", "v"),
    ("total_current_liabilities", "Total current liabilities", ("sum", "ap_accrued", "deferred_revenue")),
    ("lease_liab_noncurrent", "Operating lease liabilities, non-current", "v"),
    ("convertible_notes", "Convertible notes (fair value)", "v"),
    ("other_noncurrent_liabilities", "Other non-current liabilities", "v"),
    ("total_liabilities", "Total liabilities", ("sum", "total_current_liabilities", "other_noncurrent_liabilities")),
    None,
    ("redeemable_preferred", "Redeemable convertible preferred stock", "v"),
    None,
    ("apic", "Additional paid-in capital", "v"),
    ("aoci", "Accumulated other comprehensive income (loss)", "v"),
    ("accumulated_deficit", "Accumulated deficit", "v"),
    ("total_stockholders_deficit", "Total stockholders' equity (deficit)", ("sum", "apic", "accumulated_deficit")),
    None,
    ("total_liab_pref_eq", "Total liabilities, preferred & equity", ("add", "total_liabilities", "redeemable_preferred", "total_stockholders_deficit")),
    ("check", "Check: assets − liabilities − preferred − equity (should be 0)", ("chk",)),
]


def build_hist_bs(wb):
    ws = wb.create_sheet("Hist BS")
    d = HIST["balance_sheet"]
    title(ws, "Historical balance sheet", "S-1 consolidated balance sheets")
    hr = 4
    header_row(ws, hr, d["periods"])
    r = hr + 1
    rowmap = {}
    for item in BS_LAYOUT:
        if item is None:
            r += 1
            continue
        key, lab, typ = item
        rowmap[key] = r
        label(ws, r, lab, bold=typ != "v", italic=key == "check")
        r += 1
    for j in range(len(d["periods"])):
        col = 2 + j
        c = COL(col)
        R = lambda k: f"{c}{rowmap[k]}"
        for item in BS_LAYOUT:
            if item is None:
                continue
            key, _, typ = item
            cell = ws.cell(rowmap[key], col)
            if typ == "v":
                cell.value = d[key][j]
                style_val(cell, "s1")
                continue
            if typ[0] == "sum":
                cell.value = f"=SUM({R(typ[1])}:{R(typ[2])})"
            elif typ[0] == "add":
                cell.value = "=" + "+".join(R(k) for k in typ[1:])
            else:
                cell.value = (f"={R('total_assets')}-{R('total_liabilities')}-{R('redeemable_preferred')}"
                              f"-{R('total_stockholders_deficit')}")
            style_val(cell, bold=key != "check", italic=key == "check")
            if key != "check":
                cell.border = B_TOP
    widths(ws, 58, 13, 3)
    ws.freeze_panes = ws.cell(hr + 1, 2)
    return rowmap


# ---------------------------------------------------------------- Hist CF
CF_LAYOUT = [
    ("h", "Operating activities"),
    ("net_income", "Net income (loss)"),
    ("sbc", "Stock-based compensation"),
    ("d_and_a", "Depreciation & amortization"),
    ("fv_change_notes", "Change in fair value of convertible notes & SAFEs"),
    ("donation_matching", "Non-cash donation matching"),
    ("other_fin_instruments", "Other financial instruments"),
    ("other_noncash", "Other non-cash items"),
    ("chg_ar", "Change in accounts receivable"),
    ("chg_prepaid_other", "Change in prepaid & other assets"),
    ("chg_ap_accrued", "Change in accounts payable & accrued"),
    ("chg_leases", "Change in lease balances"),
    ("chg_deferred_rev", "Change in deferred revenue"),
    ("cfo", "Cash from operations", ("net_income", "chg_deferred_rev")),
    ("h", "Investing activities"),
    ("capex", "Capital expenditures"),
    ("purchase_st_inv", "Purchases of short-term investments"),
    ("sale_st_inv", "Sales / maturities of short-term investments"),
    ("other_investing", "Other investing"),
    ("cfi", "Cash from investing", ("capex", "other_investing")),
    ("h", "Financing activities"),
    ("pref_stock_proceeds", "Preferred stock proceeds"),
    ("finance_lease_principal", "Finance lease principal"),
    ("convertible_notes_issued", "Convertible notes issued"),
    ("common_stock_proceeds", "Common stock proceeds (option exercises)"),
    ("repurchases", "Repurchases"),
    ("safes_issued", "SAFEs issued"),
    ("offering_costs", "Offering costs"),
    ("cff", "Cash from financing", ("pref_stock_proceeds", "offering_costs")),
    ("h", ""),
    ("net_change", "Net change in cash (CFO + CFI + CFF)", None),
    ("fcf", "Free cash flow (CFO + capex)", None),
]


def build_hist_cf(wb):
    ws = wb.create_sheet("Hist CF")
    d = HIST["cash_flow"]
    periods, flags = d["periods"], d["flags"]
    title(ws, "Historical cash flow statement", "S-1 consolidated statements of cash flows; H2'25 derived as FY25 − H1'25")
    hr = 4
    header_row(ws, hr, periods)
    r = hr + 1
    rowmap = {}
    for item in CF_LAYOUT:
        if item[0] == "h":
            if item[1]:
                label(ws, r, item[1], bold=True)
            r += 1
            continue
        rowmap[item[0]] = r
        tot = len(item) == 3
        label(ws, r, item[1], bold=tot, italic=item[0] == "fcf", indent=0 if tot else 1)
        r += 1
    i_fy25, i_h125 = periods.index("FY25"), periods.index("H1'25")
    for j, p in enumerate(periods):
        col = 2 + j
        c = COL(col)
        R = lambda k: f"{c}{rowmap[k]}"
        for item in CF_LAYOUT:
            if item[0] == "h":
                continue
            key = item[0]
            cell = ws.cell(rowmap[key], col)
            subtotal = len(item) == 3
            if subtotal and item[2]:
                cell.value = f"=SUM({R(item[2][0])}:{R(item[2][1])})"
            elif key == "net_change":
                cell.value = f"={R('cfo')}+{R('cfi')}+{R('cff')}"
            elif key == "fcf":
                cell.value = f"={R('cfo')}+{R('capex')}"
            elif flags[j] == "derived":
                cell.value = f"={COL(2 + i_fy25)}{rowmap[key]}-{COL(2 + i_h125)}{rowmap[key]}"
            else:
                cell.value = d[key][j]
            style_val(cell, flags[j], bold=subtotal or key == "net_change", italic=key == "fcf")
            if subtotal:
                cell.border = B_TOP
    widths(ws, 50, 11, len(periods))
    ws.freeze_panes = ws.cell(hr + 1, 2)
    return rowmap


# ---------------------------------------------------------------- scenario sheets
RATIOS = [("grossMargin", "Gross margin (revenue vs serving compute)", PCT1),
          ("adjOpMargin", "Adjusted operating margin", PCT1),
          ("opMargin", "Operating margin", PCT1),
          ("computePctRev", "Compute % of revenue", PCT0),
          ("headcount", "Headcount (period end)", "#,##0"),
          ("fdShares", "Fully diluted shares (millions)", NUM1)]


def build_scenario(wb, key):
    sheet, name = SCEN[key]
    s = FC["scenarios"][key]
    ws = wb.create_sheet(sheet)
    ipo = s["ipo"]
    title(ws, f"{name} — forecast", f"Engine label: {s['label']}. Values from model/engine.js (data/forecast.json); not formulas. $ millions.")
    ws.cell(3, 1, f"IPO assumed {ipo['quarter']} at ${ipo['price']:.2f} per share; primary raise ${ipo['primary']:,.0f}M; "
                  f"Google placement ${ipo['googlePlacement']:,.0f}M; notes value at IPO ${ipo['notesValueAtIpo']:,.0f}M.").font = F_BOLD
    qn = 6
    qlabels = s["periods"][:qn]
    years = sorted(s["annual"].keys())
    hdrs = qlabels + [f"FY{y}" for y in years]
    ws.cell(4, 1, "FY2026 includes S-1 H1'26 actuals + Q3'26 calibration + Q4'26 forecast. Q3'26 is the calibration quarter.").font = F_NOTE

    def get(k, j):
        if j < qn:
            return s["series"][k][j]
        return s["annual"][years[j - qn]][k]

    rowmap = {}
    r = 6
    for block, blabel in [("income", "Income statement"), ("balance", "Balance sheet (period end)"), ("cashflow", "Cash flow")]:
        ws.cell(r, 1, blabel).font = F_SUB
        r += 1
        header_row(ws, r, hdrs, fill=FILL_FC)
        r += 1
        for row in FC["layout"][block]:
            k, sign, kind = row["k"], row.get("sign", 1), row.get("kind")
            label(ws, r, row["label"], bold=kind == "total", italic=kind == "memo")
            rowmap[(block, k)] = r
            for j in range(len(hdrs)):
                v = get(k, j)
                cell = ws.cell(r, 2 + j, None if v is None else v * sign)
                style_val(cell, bold=kind == "total", italic=kind == "memo")
                if kind == "total":
                    cell.border = B_TOP
            r += 1
        r += 1
    ws.cell(r, 1, "Key ratios").font = F_SUB
    r += 1
    header_row(ws, r, hdrs, fill=FILL_FC, first="")
    r += 1
    for k, lab, fmt in RATIOS:
        label(ws, r, lab)
        rowmap[("ratio", k)] = r
        for j in range(len(hdrs)):
            style_val(ws.cell(r, 2 + j, get(k, j)), fmt=fmt)
        r += 1
    widths(ws, 50, 11, len(hdrs))
    ws.freeze_panes = "B6"
    return rowmap, hdrs


# ---------------------------------------------------------------- Compute staircase
def contract_labels():
    labs = {c["id"]: c["label"] for c in ASSUME["common"]["contracts"]["cloud"]}
    labs["future"] = "Future commitments (model step-adds)"
    return labs


def build_staircase(wb):
    ws = wb.create_sheet("Compute Staircase")
    title(ws, "Compute staircase — quarterly compute cost by contract", "Engine output; $ millions. Total compute = SUM of rows (formula); check vs engine `compute` should be 0.")
    labs = contract_labels()
    r = 4
    checks = {}
    for key, (sheet, name) in SCEN.items():
        s = FC["scenarios"][key]
        periods = s["periods"]
        ser = s["series"]
        ws.cell(r, 1, name).font = F_SUB
        r += 1
        hdr = r
        header_row(ws, r, periods, fill=FILL_FC)
        r += 1
        comp_keys = [k for k in ser if k.startswith("contract_")]
        order = [f"contract_{c['id']}" for c in ASSUME["common"]["contracts"]["cloud"]] + ["contract_future"]
        comp_keys = [k for k in order if k in comp_keys] + [k for k in comp_keys if k not in order]
        rows = [(k, labs.get(k[9:], k[9:])) for k in comp_keys] + [
            ("leaseDepreciation", "Broadcom TPU finance-lease depreciation"),
            ("flexCompute", "Flexible / spot compute"),
            ("ownDepreciation", "Own equipment depreciation"),
            ("capacityResale", "Less: idle capacity resold (management response)")]
        first = r
        for k, lab in rows:
            label(ws, r, lab, indent=1)
            sign = -1 if k == "capacityResale" else 1
            for j, v in enumerate(ser[k]):
                style_val(ws.cell(r, 2 + j, sign * v))
            r += 1
        last = r - 1
        tot = r
        label(ws, r, "Total compute (sum)", bold=True)
        for j in range(len(periods)):
            c = COL(2 + j)
            cell = ws.cell(r, 2 + j, f"=SUM({c}{first}:{c}{last})")
            style_val(cell, bold=True)
            cell.border = B_TOP
        r += 1
        label(ws, r, "Engine compute (forecast.json)", italic=True)
        for j, v in enumerate(ser["compute"]):
            style_val(ws.cell(r, 2 + j, v), italic=True)
        r += 1
        label(ws, r, "Check: sum − engine (should be 0)", italic=True)
        for j in range(len(periods)):
            c = COL(2 + j)
            style_val(ws.cell(r, 2 + j, f"={c}{tot}-{c}{tot + 1}"), italic=True)
        checks[key] = r
        r += 1
        rev = r
        label(ws, r, "Revenue", bold=True)
        for j, v in enumerate(ser["revenue"]):
            style_val(ws.cell(r, 2 + j, v), bold=True)
        r += 2

        # chart: stacked components + revenue line
        bar = BarChart()
        bar.type = "col"
        bar.grouping = "stacked"
        bar.overlap = 100
        bar.gapWidth = 50
        bar.title = f"{name}: quarterly compute by contract vs revenue ($M)"
        bar.y_axis.title = "$ millions"
        bar.y_axis.numFmt = "#,##0"
        bar.y_axis.majorGridlines = None
        data = Reference(ws, min_col=1, max_col=1 + len(periods), min_row=first, max_row=last)
        bar.add_data(data, from_rows=True, titles_from_data=True)
        cats = Reference(ws, min_col=2, max_col=1 + len(periods), min_row=hdr, max_row=hdr)
        bar.set_categories(cats)
        line = LineChart()
        line.add_data(Reference(ws, min_col=1, max_col=1 + len(periods), min_row=rev, max_row=rev),
                      from_rows=True, titles_from_data=True)
        line.series[0].graphicalProperties.line.width = 28000
        line.series[0].graphicalProperties.line.solidFill = "000000"
        line.series[0].smooth = False
        line.y_axis.axId = 100  # share the bar chart's value axis
        bar.x_axis.delete = False
        bar.y_axis.delete = False
        bar += line
        bar.height, bar.width = 11, 30
        bar.legend.position = "r"
        ws.add_chart(bar, f"A{r}")
        r += 24
    widths(ws, 42, 9, 18)
    ws.freeze_panes = "B4"
    return checks


# ---------------------------------------------------------------- Scenario compare
COMPARE = [  # label, (block, key), fmt
    ("Revenue", ("income", "revenue"), NUM),
    ("Compute & infrastructure", ("income", "compute"), NUM),
    ("Adjusted operating margin", ("ratio", "adjOpMargin"), PCT1),
    ("Operating income (loss)", ("income", "operatingIncome"), NUM),
    ("Net income (loss)", ("income", "netIncome"), NUM),
    ("Cash from operations", ("cashflow", "cfo"), NUM),
    ("FCF after finance lease principal", ("cashflow", "fcfAfterLeases"), NUM),
    ("Year-end cash & short-term investments", ("balance", "liquidity"), NUM),
    ("Funding gap (new equity required)", ("cashflow", "fundingGap"), NUM),
]


def build_compare(wb, scen_maps):
    ws = wb.create_sheet("Scenario Compare")
    title(ws, "Scenario comparison — annual, FY2026–FY2030", "Green cells are formulas linking to each scenario sheet's FY columns. $ millions unless %.")
    years = sorted(FC["scenarios"]["high"]["annual"].keys())
    r = 4
    anchors = {}
    for lab, ref, fmt in COMPARE:
        ws.cell(r, 1, lab).font = F_SUB
        header_row(ws, r, [f"FY{y}" for y in years], fill=FILL_FC, first=lab)
        ws.cell(r, 1).font = F_SUB
        r += 1
        anchors[lab] = (r, r + 2)
        for key, (sheet, name) in SCEN.items():
            rowmap, hdrs = scen_maps[key]
            label(ws, r, name, indent=1)
            src_row = rowmap[ref]
            for j, y in enumerate(years):
                src_col = COL(2 + hdrs.index(f"FY{y}"))
                cell = ws.cell(r, 2 + j, f"={q(sheet)}!{src_col}{src_row}")
                style_val(cell, fmt=fmt)
                cell.font = F_LINK
            r += 1
        r += 1
    widths(ws, 42, 12, len(years))
    ws.freeze_panes = "B4"

    chart_row = r + 1
    for i, (lab, fmt, ttl) in enumerate([("Revenue", "#,##0", "Revenue by scenario ($M)"),
                                         ("Adjusted operating margin", "0%", "Adjusted operating margin by scenario")]):
        lo, hi = anchors[lab]
        ch = LineChart()
        ch.title = ttl
        ch.y_axis.numFmt = fmt
        ch.add_data(Reference(ws, min_col=1, max_col=1 + len(years), min_row=lo, max_row=hi), from_rows=True, titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, max_col=1 + len(years), min_row=lo - 1, max_row=lo - 1))
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.height, ch.width = 9, 16
        ch.legend.position = "b"
        ws.add_chart(ch, f"{'A' if i == 0 else 'D'}{chart_row}")


# ---------------------------------------------------------------- Assumptions
def fmt_value(v):
    if v is None:
        return "null (none)"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return v


def flatten(d, prefix, out, skip=()):
    """Rows of (path, value, notes). Annotation keys (_src, _note, _bud, _src_X, _X_note) go to notes."""
    notes = {k: v for k, v in d.items() if k.startswith("_")}
    keys = [k for k in d if not k.startswith("_") and k not in skip]
    targeted = {}
    generic = []
    for nk, nv in notes.items():
        tgt = None
        if nk.startswith("_src_") and nk[5:] in keys:
            tgt = nk[5:]
        elif nk.endswith("_note") and nk[1:-5] in keys:
            tgt = nk[1:-5]
        if tgt:
            targeted.setdefault(tgt, []).append(nv)
        else:
            generic.append(f"[{nk[1:]}] {nv}")
    start = len(out)
    for k in keys:
        v = d[k]
        path = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict):
            flatten(v, path, out)
            if k in targeted and len(out) > 0:
                p, val, n = out[-1]
                out[-1] = (p, val, "; ".join(filter(None, [n] + targeted[k])))
        else:
            out.append((path, fmt_value(v), "; ".join(targeted.get(k, []))))
    if generic:
        if len(out) > start:
            p, val, n = out[start]
            out[start] = (p, val, "; ".join(filter(None, generic + [n])))
        else:
            out.append((prefix, "", "; ".join(generic)))


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    title(ws, "Model assumptions (model/assumptions.json)", ASSUME.get("_readme", ""))
    ws.cell(2, 1).alignment = Alignment(wrap_text=False)
    r = 4
    sections = [("Common parameters (all scenarios)", ASSUME["common"], ("contracts",))]
    for key, (_, name) in SCEN.items():
        sections.append((f"Overrides — {name}", ASSUME["scenarios"][key], ()))
    for heading, d, skip in sections:
        ws.cell(r, 1, heading).font = F_SUB
        r += 1
        for j, h in enumerate(["Parameter", "Value", "Source / note"]):
            c = ws.cell(r, 1 + j, h)
            c.font = F_HDR
            c.fill = FILL_HIST_HDR
            c.border = B_BOTTOM
        r += 1
        rows = []
        flatten(d, "", rows, skip)
        for path, val, note in rows:
            ws.cell(r, 1, path).font = F_BASE
            c = ws.cell(r, 2, val)
            c.font = F_BASE
            c.alignment = Alignment(horizontal="left")
            if isinstance(val, float) and abs(val) < 1:
                c.number_format = "0.0##"
            elif isinstance(val, (int, float)):
                c.number_format = "#,##0.###"
            n = ws.cell(r, 3, note)
            n.font = F_NOTE
            n.alignment = WRAP
            r += 1
        r += 1

    ws.cell(r, 1, "Contracts (common; scenario overrides above adjust ramps / deferrals / cancellation)").font = F_SUB
    r += 1
    hdrs = ["Contract", "Total or full rate ($M)", "Source note", "Start", "End", "Ramp (qtrs) / profile", "Type / flags"]
    for j, h in enumerate(hdrs):
        c = ws.cell(r, 1 + j, h)
        c.font = F_HDR
        c.fill = FILL_HIST_HDR
        c.border = B_BOTTOM
        c.alignment = WRAP
    r += 1
    con = ASSUME["common"]["contracts"]
    for c in con["cloud"]:
        if c.get("type") == "annualSchedule":
            amt = "Schedule: " + ", ".join(f"{y}: {v:,}" for y, v in c["schedule"].items())
            start, end = min(c["schedule"]), max(c["schedule"])
        elif "total" in c:
            amt, start, end = c["total"], c["start"], c["end"]
        else:
            amt = f"{c['fullRate']:,}/qtr full rate"
            start, end = c["start"], c["end"]
        ramp = c.get("ramp", ", ".join(str(x) for x in c.get("rampProfile", [])) or "")
        flags = ", ".join(x for x in [c.get("type", ""), "cancelable" if c.get("cancelable") else ""] if x)
        vals = [f"{c['label']} ({c['id']})", amt, c.get("_src", con.get("_src", "")), start, end, ramp, flags]
        for j, v in enumerate(vals):
            cell = ws.cell(r, 1 + j, v)
            cell.font = F_NOTE if j == 2 else F_BASE
            cell.alignment = WRAP
            if isinstance(v, (int, float)):
                cell.number_format = "#,##0"
        r += 1
    fl = con["financeLeases"]
    commence = ", ".join(f"{k}: {v:.0%}" for k, v in fl["commence"].items())
    vals = [f"{fl['label']} ({fl['id']})", fl["total"], fl.get("_src", ""), min(fl["commence"]), f"{fl['termQuarters']}-qtr terms",
            f"commencement: {commence}", f"finance lease @ {fl['rate']:.1%}"]
    for j, v in enumerate(vals):
        cell = ws.cell(r, 1 + j, v)
        cell.font = F_NOTE if j == 2 else F_BASE
        cell.alignment = WRAP
        if isinstance(v, (int, float)):
            cell.number_format = "#,##0"
    r += 1
    fcm = ASSUME["common"]["futureCommit"]
    vals = ["Future commitments (future)", f"+{fcm['stepAdd']:,}/qtr each quarter (common; see overrides)", fcm.get("_note", ""),
            fcm["start"], "horizon end", "", "model"]
    for j, v in enumerate(vals):
        cell = ws.cell(r, 1 + j, v)
        cell.font = F_NOTE if j == 2 else F_BASE
        cell.alignment = WRAP
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    for col, w in zip("DEFG", (12, 14, 22, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------- main
def main():
    wb = Workbook()
    build_readme(wb)
    build_hist_is_q(wb)
    build_hist_is_a(wb)
    build_hist_bs(wb)
    build_hist_cf(wb)
    scen_maps = {k: build_scenario(wb, k) for k in SCEN}
    build_staircase(wb)
    build_compare(wb, scen_maps)
    build_assumptions(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
